# -*- coding: utf-8 -*-
"""由整理版 Markdown 生成 Word 文档：标题样式 + 首行缩进2字符 + 引文上标 + 参考文献悬挂缩进"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import docx.opc.constants as opc_const

SRC = "脑机接口外设0820初稿_整理版.md"
DST = "脑机接口外设0820初稿_整理版.docx"

HEI = "黑体"
SONG = "宋体"

# 引文角标模式：数字、逗号、连接号（– 或 -）组成的方括号
CITE_RE = re.compile(r'\[[0-9]+(?:[–\-][0-9]+)?(?:,[0-9]+(?:[–\-][0-9]+)?)*\]')
# 内联加粗 + 引文角标 联合分词
TOKEN_RE = re.compile(r'(\*\*[^*]+\*\*|\[[0-9]+(?:[–\-][0-9]+)?(?:,[0-9]+(?:[–\-][0-9]+)?)*\])')


def set_run_font(run, east_asia, size_pt, bold=None, superscript=None, color=None):
    run.font.name = east_asia
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east_asia)
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if superscript is not None:
        run.font.superscript = superscript
    if color is not None:
        run.font.color.rgb = color


def set_first_line_indent_chars(paragraph, chars=2, font_pt=12):
    """首行缩进 N 字符（w:firstLineChars 为主，w:firstLine 为兜底绝对值）"""
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:firstLineChars'), str(int(chars * 100)))
    ind.set(qn('w:firstLine'), str(int(chars * font_pt * 20)))


def add_hyperlink(paragraph, url, text, size_pt=10.5):
    part = paragraph.part
    r_id = part.relate_to(url, opc_const.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), SONG)
    rFonts.set(qn('w:eastAsia'), SONG)
    rPr.append(rFonts)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size_pt * 2))); rPr.append(sz)
    szCs = OxmlElement('w:szCs'); szCs.set(qn('w:val'), str(int(size_pt * 2))); rPr.append(szCs)
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); rPr.append(color)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_body_paragraph(doc, text, font, size, indent_chars=None, spacing=1.5):
    """添加正文段落：解析 **加粗** 与 [引文] 角标"""
    p = doc.add_paragraph()
    if spacing:
        p.paragraph_format.line_spacing = spacing
    if indent_chars:
        set_first_line_indent_chars(p, indent_chars, size)
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            run = p.add_run(text[pos:m.start()])
            set_run_font(run, font, size)
        token = m.group(0)
        if token.startswith('**'):
            run = p.add_run(token[2:-2])
            set_run_font(run, font, size, bold=True)
        else:
            # 引文角标：上标
            run = p.add_run(token)
            set_run_font(run, font, size, superscript=True)
        pos = m.end()
    if pos < len(text):
        run = p.add_run(text[pos:])
        set_run_font(run, font, size)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(text)
        set_run_font(run, HEI, 16, bold=True)
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(text)
        set_run_font(run, HEI, 14, bold=True)
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        set_run_font(run, HEI, 12, bold=True)
    return p


def add_third_heading(doc, text):
    """三级标题：粗体 1.1 标题"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, HEI, 12, bold=True)
    return p


doc = Document()

# 默认样式：正文宋体
normal = doc.styles['Normal']
normal.font.name = SONG
normal.font.size = Pt(12)
normal._element.rPr.rFonts.set(qn('w:eastAsia'), SONG)

lines = open(SRC, encoding='utf-8').read().split('\n')
i = 0
in_refs = False
while i < len(lines):
    line = lines[i]
    s = line.strip()

    if s == '':
        i += 1
        continue

    if s == '## 参考文献':
        in_refs = True
        add_heading(doc, '参考文献', 2)
        i += 1
        continue

    if in_refs:
        # 参考文献条目：N. [title](url)  →  [N] title，悬挂缩进
        m = re.match(r'^(\d+)\.\s\[(.*)\]\((.*)\)$', s)
        if m:
            num, title, url = m.group(1), m.group(2), m.group(3)
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(2)
            # 悬挂缩进：左缩进 + 首行负缩进
            p.paragraph_format.left_indent = Cm(1.1)
            p.paragraph_format.first_line_indent = Cm(-1.1)
            r1 = p.add_run(f"[{num}] ")
            set_run_font(r1, SONG, 10.5)
            add_hyperlink(p, url, title, size_pt=10.5)
        i += 1
        continue

    if s.startswith('### '):
        add_heading(doc, s[4:], 3)
    elif s.startswith('## '):
        add_heading(doc, s[3:], 2)
    elif s.startswith('# '):
        add_heading(doc, s[2:], 1)
    elif s.startswith('**') and s.endswith('**') and s.count('**') == 2:
        # 三级标题（整行粗体）
        add_third_heading(doc, s[2:-2])
    else:
        # 正文段落
        add_body_paragraph(doc, s, SONG, 12, indent_chars=2, spacing=1.5)
    i += 1

doc.save(DST)
print("已生成:", DST)
